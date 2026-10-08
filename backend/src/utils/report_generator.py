#!/usr/bin/env python3
"""Rahbar uchun mukammal Excel hisobot generatori.

To'rt asosiy varaq:
  «Umumiy»         — Asosiy ko'rsatkichlar kartalari, vazifalar holati va xodimlar darajasi (grafiklar bilan);
  «Xodimlar»       — Har bir xodim samaradorligi, bajarilgan vazifalar, muddatga rioya, daraja va databarlar;
  «Vazifalar»      — Barcha vazifalar reyestri: holatlar ranglari, to'g'ridan-to'g'ri katakka joylangan
                     FOTOSURAT dalillari miniatyuralari (thumbnails), audio va hujjatlarga faol giperhavolalar;
  «Ijro dalillari» — Har bir topshiriq bo'yicha "Rahbar topshirig'i" va "Xodim ijrosi (dalillar)" yonma-yon
                     ko'rinadigan foto-dosye kartalari.
"""
import sys
import os
import json
import math
import base64
import uuid
import shutil
import tempfile
from datetime import datetime, date
from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.chart import BarChart, DoughnutChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.chart.series import DataPoint
from openpyxl.chart.text import RichText
from openpyxl.drawing.text import CharacterProperties, Paragraph, ParagraphProperties
from openpyxl.drawing.image import Image as OpenpyxlImage
from openpyxl.formatting.rule import DataBarRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.properties import PageSetupProperties
from openpyxl.worksheet.worksheet import Worksheet
from PIL import Image as PILImage, ImageOps

DATE_FMT = "DD.MM.YYYY"
DATETIME_FMT = "DD.MM.YYYY HH:MM"
PERCENT_FMT = "0%"
DAYS_FMT = "0.0"
DAYS_UNIT_FMT = '0.0" kun"'

ON_TIME = "Muddatida"
LATE = "Kechikib"
OVERDUE = "Muddati o'tgan"
NOT_DUE_YET = "Muddat kelmagan"
NO_DUE = "Muddatsiz"
NOT_APPLICABLE = "—"
NOT_RATED = "Baholanmagan"
NO_TASKS = "Vazifa yo'q"

LEVELS = ((90, "A'lo", "🏆"), (70, "Yaxshi", "👍"), (50, "Qoniqarli", "🙂"), (0, "Past", "⚠️"))

NAVY, INK, MUTED, LINE, ZEBRA, WHITE = "1F3864", "262626", "7F7F7F", "D9D9D9", "F5F8FC", "FFFFFF"
LINK_BLUE = "0563C1"

TONES = {
    "green": ("E2EFDA", "375623"),
    "blue": ("DDEBF7", "1F4E78"),
    "amber": ("FFF2CC", "806000"),
    "red": ("FCE4E4", "9C0006"),
    "grey": ("EDEDED", "595959"),
    "purple": ("EAE4F2", "5B2C83"),
}

SOLID = {
    "green": "70AD47",
    "blue": "5B9BD5",
    "amber": "FFC000",
    "red": "E15759",
    "grey": "A5A5A5",
    "purple": "8E6BBF",
}

LEVEL_TONES = {
    "A'lo": "green",
    "Yaxshi": "blue",
    "Qoniqarli": "amber",
    "Past": "red",
    NOT_RATED: "grey",
    NO_TASKS: "grey",
}

# Talab etilgan ranglar: Kutilmoqda — qizil, Ko'rilgan — kulrang, Tasdiq — sariq, Bajarilgan — yashil
STATUS_LABELS = {
    "pending": "Kutilmoqda",
    "in_progress": "Xodim ko'rgan",
    "submitted": "Tasdiq kutilmoqda",
    "completed": "Bajarilgan",
}

STATUS_TONES = {
    "Bajarilgan": "green",
    "Bajarildi": "green",
    "Kutilmoqda": "red",
    "Xodim ko'rgan": "grey",
    "Ko'rildi": "grey",
    "Tasdiq kutilmoqda": "amber",
    "Jarayonda": "amber",
    "Bekor qilingan": "grey",
}

DEADLINE_TONES = {
    ON_TIME: "green",
    LATE: "amber",
    OVERDUE: "red",
    NOT_DUE_YET: "blue",
}

MEDALS = ("FFD966", "D9D9D9", "F4B183")

CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
TEXT = Alignment(horizontal="left", vertical="center", wrap_text=True)
HEADER_FONT = Font(size=10, bold=True, color=WHITE)
BODY_FONT = Font(size=10, color=INK)

WIDE, NARROW = 13, 2
GAPS = (4, 7, 10)
CARDS = ((2, 3), (5, 6), (8, 9), (11, 12))
LEFT, RIGHT = (2, 6), (8, 12)
CHART_DATA_COL = 14
ROW_H = 20
MAX_ROW_H = 300
PX_PER_CM, CM_PER_PT = 37.8, 2.54 / 72
TOP_CHART = 10

def _fill(color: str) -> PatternFill:
    return PatternFill("solid", fgColor=color)

def _side(color: str = LINE, style: str = "thin") -> Side:
    return Side(style=style, color=color)

GRID = Border(left=_side(), right=_side(), top=_side(), bottom=_side())
UNDERLINE = Border(bottom=_side())

FORMULA_PREFIXES = ("=", "+", "-", "@", "\t", "\r")

def _clean(value):
    if not isinstance(value, str):
        return value
    cleaned = ILLEGAL_CHARACTERS_RE.sub("", value)
    stripped = cleaned.lstrip()
    if stripped and stripped.startswith(FORMULA_PREFIXES):
        return f"'{cleaned}"
    return cleaned

def _set(ws: Worksheet, row: int, col: int, value, fmt: str | None = None):
    clean_val = _clean(value)
    cell = ws.cell(row=row, column=col, value=clean_val)
    if isinstance(clean_val, str) and (clean_val.startswith(FORMULA_PREFIXES) or clean_val.startswith("'")):
        cell.data_type = "s"
    if fmt:
        cell.number_format = fmt
    elif isinstance(value, datetime):
        cell.number_format = DATETIME_FMT
    elif isinstance(value, date):
        cell.number_format = DATE_FMT
    return cell

def _or_dash(value):
    return NOT_APPLICABLE if value is None else value

def _tone(cell, tone: str, bold: bool = False) -> None:
    fill, color = TONES[tone]
    cell.fill = _fill(fill)
    cell.font = Font(size=cell.font.size, bold=bold, color=color)

def _block(ws: Worksheet, row: int, c1: int, c2: int, value, font: Font, fill: PatternFill | None = None,
           align: Alignment = CENTER, border: Border | None = None, fmt: str | None = None):
    if c2 > c1:
        ws.merge_cells(start_row=row, start_column=c1, end_row=row, end_column=c2)
    cell = _set(ws, row, c1, value, fmt)
    cell.font, cell.alignment = font, align
    for col in range(c1, c2 + 1):
        part = ws.cell(row=row, column=col)
        if fill:
            part.fill = fill
        if border:
            part.border = border
    return cell

def _section(ws: Worksheet, row: int, c1: int, c2: int, title: str) -> None:
    _block(ws, row, c1, c2, title.upper(), Font(size=11, bold=True, color=NAVY),
           align=Alignment(horizontal="left", vertical="bottom"), border=Border(bottom=_side(NAVY, "medium")))
    ws.row_dimensions[row].height = 24

def _print_setup(ws: Worksheet, landscape: bool) -> None:
    ws.page_setup.orientation = "landscape" if landscape else "portrait"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)

def _labels(**shown) -> DataLabelList:
    flags = dict(showVal=False, showPercent=False, showCatName=False, showSerName=False,
                 showLegendKey=False, showLeaderLines=False)
    return DataLabelList(**(flags | shown))

def _white_bold() -> RichText:
    props = CharacterProperties(b=True, sz=1000, solidFill=WHITE)
    return RichText(p=[Paragraph(pPr=ParagraphProperties(defRPr=props), endParaRPr=props)])

def _chart_data(ws: Worksheet, chart, parts: list[tuple[str, float, str]], row: int, col: int):
    for i, (label, value, _) in enumerate(parts):
        _set(ws, row + i, col, label)
        _set(ws, row + i, col + 1, value)
    last = row + len(parts) - 1
    chart.add_data(Reference(ws, min_col=col + 1, min_row=row, max_row=last))
    chart.set_categories(Reference(ws, min_col=col, min_row=row, max_row=last))
    chart.visible_cells_only = False
    for c in (col, col + 1):
        ws.column_dimensions[get_column_letter(c)].hidden = True
    series = chart.series[0]
    for i, (_, _, color) in enumerate(parts):
        point = DataPoint(idx=i)
        point.graphicalProperties.solidFill = color
        point.graphicalProperties.line.solidFill = WHITE
        series.dPt.append(point)
    return series

def _place(ws: Worksheet, chart, row: int, rows: int) -> None:
    width_px = sum((NARROW if c in GAPS else WIDE) * 7 + 5 for c in range(RIGHT[0], RIGHT[1] + 1))
    chart.width = width_px / PX_PER_CM - 0.1
    chart.height = rows * ROW_H * CM_PER_PT
    ws.add_chart(chart, f"{get_column_letter(RIGHT[0])}{row}")

def parse_date(val):
    if not val:
        return None
    if isinstance(val, (datetime, date)):
        return val
    try:
        val_str = str(val).strip()
        if "T" in val_str:
            return datetime.fromisoformat(val_str.replace("Z", "+00:00")).replace(tzinfo=None)
        if len(val_str) >= 19 and " " in val_str:
            return datetime.strptime(val_str[:19], "%Y-%m-%d %H:%M:%S")
        return datetime.strptime(val_str[:10], "%Y-%m-%d")
    except Exception:
        return None

def compute_task_state(task, now):
    status = task.get("status")
    due_date = parse_date(task.get("due_date"))
    updated_at = parse_date(task.get("updated_at"))

    if not due_date:
        return NO_DUE

    if status == "completed":
        if updated_at:
            due_end = datetime(due_date.year, due_date.month, due_date.day, 23, 59, 59)
            return ON_TIME if updated_at <= due_end else LATE
        return ON_TIME

    due_end = datetime(due_date.year, due_date.month, due_date.day, 23, 59, 59)
    return OVERDUE if due_end < now else NOT_DUE_YET

class UserStats:
    def __init__(self):
        self.total = 0
        self.completed = 0
        self.on_time = 0
        self.late = 0
        self.review = 0       # pending / kutilmoqda
        self.in_work = 0      # in_progress / xodim ko'rgan / submitted
        self.overdue = 0
        self.cancelled = 0
        self.days = []

    def add(self, task, now):
        self.total += 1
        st = task.get("status")
        state = compute_task_state(task, now)
        created_at = parse_date(task.get("created_at"))
        updated_at = parse_date(task.get("updated_at"))

        if st == "completed":
            self.completed += 1
            if state == LATE:
                self.late += 1
            else:
                self.on_time += 1
            if updated_at and created_at:
                diff = max(0.0, (updated_at - created_at).total_seconds() / 86400)
                self.days.append(diff)
        elif st == "pending":
            self.review += 1
            if state == OVERDUE:
                self.overdue += 1
        elif st in ("in_progress", "submitted"):
            self.in_work += 1
            if state == OVERDUE:
                self.overdue += 1

    @property
    def completion(self):
        base = self.total - self.cancelled
        return self.completed / base if base else None

    @property
    def score(self):
        decided = self.completed + self.overdue
        if not decided:
            return None
        on_time = self.completed - self.late
        return round(100 * (on_time + 0.5 * self.late) / decided)

    @property
    def avg_days(self):
        return round(sum(self.days) / len(self.days), 1) if self.days else None

    @property
    def level(self):
        if self.score is None:
            return NOT_RATED if self.total else NO_TASKS
        for threshold, name, _ in LEVELS:
            if self.score >= threshold:
                return name
        return "Past"

GLOSSARY = (
    ("Kutilmoqda", "Yangi yaratilgan, ijrochi tomonidan ko'rib chiqilishi kutilayotgan topshiriqlar (qizil)."),
    ("Xodim ko'rgan", "Xodim tomonidan tanishilgan va ayni damda bajarilayotgan topshiriqlar (kulrang)."),
    ("Tasdiq kutilmoqda", "Xodim hisobot va ijro dalillarini yuklab rahbar tasdig'iga yuborgan topshiriqlar (sariq)."),
    ("Bajarilgan", "Rahbar tomonidan tasdiqlangan va to'liq yakunlangan topshiriqlar (yashil)."),
    ("Ijro dalillari", "Xodim hisoboti bilan biriktirilgan fotosuratlar, audio hisobotlar, dalolatnoma va smetalar."),
    ("Samaradorlik bali", "0–100 ball: muddatida bajarilgan (100%), kechikib (50%), muddati o'tgan (0%)."),
)
GLOSSARY_CHARS = 90

def _cards(ws: Worksheet, total: UserStats, users: list) -> None:
    avg = f"{total.avg_days:.1f}".replace(".", ",") if total.avg_days is not None else None
    cards = (
        ("Jami vazifalar", total.total, None, f"xodimlar: {len(users)} ta", "blue"),
        ("Bajarilgan", total.completed, None, f"muddatida {total.on_time} · kechikib {total.late}", "green"),
        ("Jarayonda", total.in_work + total.review, None, f"shundan muddati o'tgan: {total.overdue}",
         "red" if total.overdue else "amber"),
        ("Bajarilish %", _or_dash(total.completion), PERCENT_FMT,
         f"o'rtacha bajarish: {avg} kun" if avg else "hali bajarilgani yo'q", "purple"),
    )
    for (c1, c2), (label, value, fmt, note, tone) in zip(CARDS, cards):
        fill, color = _fill(TONES[tone][0]), TONES[tone][1]
        _block(ws, 4, c1, c2, label.upper(), Font(size=9, bold=True, color=color), fill,
               border=Border(top=_side(SOLID[tone], "thick")))
        _block(ws, 5, c1, c2, value, Font(size=24, bold=True, color=color), fill, fmt=fmt)
        _block(ws, 6, c1, c2, note, Font(size=9, color=color), fill)

def _status_section(ws: Worksheet, total: UserStats, row: int) -> int:
    (l1, l2), (r1, r2) = LEFT, RIGHT
    _section(ws, row, l1, l2, "Vazifalar holati")
    lines = (
        ("Jami vazifalar", total.total, NAVY, None, False),
        ("Bajarilgan", total.completed, SOLID["green"], None, False),
        ("muddatida", total.on_time, None, None, False),
        ("kechikib", total.late, None, None, False),
        ("Kutilmoqda", total.review, SOLID["red"], None, False),
        ("Jarayonda (ko'rilgan)", total.in_work, SOLID["grey"], None, False),
        ("shundan muddati o'tgan", total.overdue, None, None, total.overdue > 0),
        ("Bekor qilingan", total.cancelled, SOLID["grey"], None, False),
        ("Bajarilish %", _or_dash(total.completion), SOLID["purple"], PERCENT_FMT, False),
        ("O'rtacha bajarish vaqti", _or_dash(total.avg_days), SOLID["purple"], DAYS_UNIT_FMT, False),
    )
    for i, (label, value, marker, fmt, alert) in enumerate(lines, row + 1):
        sub = marker is None
        color = MUTED if sub else INK
        ws.row_dimensions[i].height = ROW_H
        _block(ws, i, l1, l2 - 1, label, Font(size=10, bold=not sub, color=color),
               align=Alignment(vertical="center", indent=3 if sub else 1), border=UNDERLINE)
        if marker:
            ws.cell(row=i, column=l1).border = Border(left=_side(marker, "thick"), bottom=_side())
        cell = _block(ws, i, l2, l2, value, Font(size=10 if sub else 11, bold=not sub, color=color),
                      border=UNDERLINE, fmt=fmt)
        if alert:
            _tone(cell, "red", bold=True)

    parts = [(name, value, SOLID[tone]) for name, value, tone in (
        ("Bajarilgan", total.completed, "green"),
        ("Kutilmoqda", total.review, "red"),
        ("Jarayonda", max(0, total.in_work - total.overdue), "blue"),
        ("Muddati o'tgan", total.overdue, "red"),
        ("Bekor qilingan", total.cancelled, "grey"),
    ) if value]

    if parts:
        _section(ws, row, r1, r2, "Holatlar ulushi")
        chart = DoughnutChart(holeSize=55)
        series = _chart_data(ws, chart, parts, row + 1, CHART_DATA_COL)
        series.dLbls = _labels(showPercent=True)
        series.dLbls.txPr = _white_bold()
        chart.legend.position = "r"
        _place(ws, chart, row + 1, len(lines))
    return row + len(lines) + 1

def _levels_section(ws: Worksheet, ranked: list, row: int) -> int:
    (l1, l2), (r1, r2) = LEFT, RIGHT
    _section(ws, row, l1, l2, "Xodimlar darajasi")
    head = row + 1
    ws.row_dimensions[head].height = ROW_H
    header_font, header_fill = Font(size=9, bold=True, color=MUTED), _fill("F2F2F2")
    for c1, c2, title in ((l1, l1 + 1, "Daraja"), (l1 + 2, l2 - 1, "Samaradorlik bali"), (l2, l2, "Xodimlar")):
        _block(ws, head, c1, c2, title, header_font, header_fill, border=UNDERLINE)

    spans = [f"{level[0]}–{LEVELS[i - 1][0] - 1 if i else 100}" for i, level in enumerate(LEVELS)]
    levels = [(name, span) for (_, name, _), span in zip(LEVELS, spans)]
    levels += [(NOT_RATED, NOT_APPLICABLE), (NO_TASKS, NOT_APPLICABLE)]
    for i, (name, span) in enumerate(levels, head + 1):
        ws.row_dimensions[i].height = ROW_H
        fill, color = TONES[LEVEL_TONES[name]]
        _block(ws, i, l1, l1 + 1, name, Font(size=10, bold=True, color=color), _fill(fill),
               align=Alignment(vertical="center", indent=1), border=UNDERLINE)
        _block(ws, i, l1 + 2, l2 - 1, span, Font(size=10, color=MUTED), border=UNDERLINE)
        count = sum(s.level == name for _, s in ranked)
        _block(ws, i, l2, l2, count, Font(size=11, bold=True, color=INK if count else MUTED), border=UNDERLINE)
    end = head + len(levels)

    rated = [(u, s) for u, s in ranked if s.score is not None]
    if rated:
        title = "Samaradorlik bali" + (f" · eng yaxshi {TOP_CHART}" if len(rated) > TOP_CHART else "")
        rated = rated[:TOP_CHART]
        _section(ws, row, r1, r2, title)
        chart = BarChart()
        chart.type, chart.gapWidth, chart.legend = "bar", 60, None
        parts = [(u.get("name", "Xodim"), s.score, SOLID[LEVEL_TONES[s.level]]) for u, s in rated]
        series = _chart_data(ws, chart, parts, head, CHART_DATA_COL + 3)
        series.dLbls = _labels(showVal=True)
        series.dLbls.dLblPos = "outEnd"
        chart.x_axis.scaling.orientation = "maxMin"
        chart.x_axis.delete = False
        chart.y_axis.delete = True
        chart.y_axis.majorGridlines = None
        chart.y_axis.scaling.min, chart.y_axis.scaling.max = 0, 115
        rows = max(len(levels) + 1, math.ceil((1.2 + 0.75 * len(rated)) / (ROW_H * CM_PER_PT)))
        _place(ws, chart, head, rows)
        end = max(end, head + rows - 1)
    return end + 1

def _glossary(ws: Worksheet, row: int) -> None:
    _section(ws, row, LEFT[0], RIGHT[1], "Izoh")
    for i, (term, text) in enumerate(GLOSSARY, row + 1):
        ws.row_dimensions[i].height = max(ROW_H, 14 * math.ceil(len(text) / GLOSSARY_CHARS) + 6)
        _block(ws, i, LEFT[0], LEFT[0] + 1, term, Font(size=10, bold=True, color=NAVY),
               align=Alignment(vertical="center", wrap_text=True), border=UNDERLINE)
        _block(ws, i, LEFT[0] + 2, RIGHT[1], text, BODY_FONT, align=TEXT, border=UNDERLINE)

def _summary_sheet(ws: Worksheet, total: UserStats, ranked: list, users: list, author: dict, now: datetime) -> None:
    ws.title = "Umumiy"
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = NAVY
    ws.column_dimensions["A"].width = NARROW
    for col in range(2, RIGHT[1] + 1):
        ws.column_dimensions[get_column_letter(col)].width = NARROW if col in GAPS else WIDE

    banner = _fill(NAVY)
    _block(ws, 1, 2, RIGHT[1], "Xodimlar faoliyati bo'yicha hisobot", Font(size=18, bold=True, color=WHITE),
           banner, Alignment(vertical="center", indent=1))
    who = author.get("name") if author else NOT_APPLICABLE
    _block(ws, 2, 2, RIGHT[1], f"Tuzilgan: {now:%d.%m.%Y %H:%M}      Tuzuvchi: {who}",
           Font(size=10, color="D6DCE5"), banner, Alignment(vertical="top", indent=1))
    for r, height in ((1, 36), (2, 22), (3, 12), (4, 22), (5, 42), (6, 22), (7, 12)):
        ws.row_dimensions[r].height = height

    _cards(ws, total, users)
    row = _status_section(ws, total, 8)
    row = _levels_section(ws, ranked, row + 1)
    _glossary(ws, row + 1)
    _print_setup(ws, landscape=False)

class Column:
    def __init__(self, title, width, fmt=None, text=False):
        self.title = title
        self.width = width
        self.fmt = fmt
        self.text = text

def _row_height(columns: list, values: list) -> float:
    lines = 1
    for column, value in zip(columns, values):
        if column.text and isinstance(value, str) and value:
            per_line = max(1, int(column.width * 1.1))
            lines = max(lines, sum(math.ceil(max(len(part), 1) / per_line) for part in value.split("\n")))
    return min(MAX_ROW_H, max(ROW_H, 13 * lines + 7))

def _data_sheet(ws: Worksheet, columns: list, rows: list, tab_color: str) -> dict:
    ws.sheet_properties.tabColor = tab_color
    header_fill = _fill(NAVY)
    for col, column in enumerate(columns, 1):
        cell = _set(ws, 1, col, column.title)
        cell.font, cell.fill, cell.alignment, cell.border = HEADER_FONT, header_fill, CENTER, GRID
        ws.column_dimensions[get_column_letter(col)].width = column.width
    ws.row_dimensions[1].height = 36
    zebra = _fill(ZEBRA)
    for r, values in enumerate(rows, 2):
        for col, (column, value) in enumerate(zip(columns, values), 1):
            cell = _set(ws, r, col, value, column.fmt)
            cell.font, cell.border = BODY_FONT, GRID
            cell.alignment = TEXT if column.text else CENTER
            if r % 2:
                cell.fill = zebra
        ws.row_dimensions[r].height = _row_height(columns, values)
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(columns))}{len(rows) + 1}"
    ws.print_title_rows = "1:1"
    _print_setup(ws, landscape=True)
    return {column.title: col for col, column in enumerate(columns, 1)}

def _data_bar(ws: Worksheet, col: int, last_row: int, maximum: float, color: str) -> None:
    letter = get_column_letter(col)
    ws.conditional_formatting.add(f"{letter}2:{letter}{last_row}", DataBarRule(
        start_type="num", start_value=0, end_type="num", end_value=maximum, color=color))

def _employees_sheet(ws: Worksheet, ranked: list) -> None:
    columns = [
        Column("O'rin", 6), Column("Xodim", 26, text=True), Column("Hudud", 22, text=True),
        Column("Lavozim", 20, text=True), Column("Holat", 15),
        Column("Jami vazifa", 9), Column("Bajarilgan", 11), Column("Muddatida", 10), Column("Kechikib", 9),
        Column("Kutilmoqda", 11), Column("Jarayonda", 10), Column("Muddati o'tgan", 10),
        Column("Bekor qilingan", 10), Column("Bajarilish %", 12, PERCENT_FMT),
        Column("O'rtacha bajarish, kun", 11, DAYS_FMT), Column("Samaradorlik bali", 13), Column("Daraja", 14),
    ]
    rows = [
        [i, u.get("name", ""), u.get("district") or u.get("region") or "",
         u.get("position") or (u.get("role") == "admin" and "Administrator" or "Xodim"),
         "faol", s.total, s.completed, s.on_time, s.late, s.review, s.in_work, s.overdue, s.cancelled,
         s.completion, s.avg_days, s.score, s.level]
        for i, (u, s) in enumerate(ranked, 1)
    ]
    col = _data_sheet(ws, columns, rows, SOLID["green"])
    for r, (u, s) in enumerate(ranked, 2):
        if r - 2 < len(MEDALS) and s.score is not None:
            medal = ws.cell(row=r, column=col["O'rin"])
            medal.fill, medal.font = _fill(MEDALS[r - 2]), Font(size=10, bold=True, color=INK)
        ws.cell(row=r, column=col["Xodim"]).font = Font(size=10, bold=True, color=INK)
        ws.cell(row=r, column=col["Holat"]).font = Font(size=10, color=TONES["green"][1])
        _tone(ws.cell(row=r, column=col["Daraja"]), LEVEL_TONES[s.level], bold=True)
        if s.overdue:
            _tone(ws.cell(row=r, column=col["Muddati o'tgan"]), "red", bold=True)
        if s.late:
            _tone(ws.cell(row=r, column=col["Kechikib"]), "amber")
    if ranked:
        last = len(ranked) + 1
        _data_bar(ws, col["Bajarilish %"], last, 1, SOLID["green"])
        _data_bar(ws, col["Samaradorlik bali"], last, 100, SOLID["blue"])

# ==============================================================================
# MULTIMEDIA VA DALILLAR HELPERLARI (Rasmlar, Audio, Fayllar, Giperhavolalar)
# ==============================================================================

def _parse_items(items):
    """Media yoki biriktirilgan fayllar ro'yxatini toza dict list ko'rinishiga keltiradi."""
    if not items:
        return []
    if isinstance(items, str):
        try:
            items = json.loads(items)
        except Exception:
            items = items.strip()
            if items:
                return [{"url": items, "name": os.path.basename(items)}]
            return []
    if isinstance(items, dict):
        return [items]
    if isinstance(items, list):
        res = []
        for it in items:
            if isinstance(it, dict):
                res.append(it)
            elif isinstance(it, str) and it.strip():
                res.append({"url": it.strip(), "name": os.path.basename(it.strip())})
        return res
    return []

def _is_image(name_or_type, url=""):
    s = f"{name_or_type or ''} {url or ''}".lower()
    return any(s.endswith(ext) for ext in (".png", ".jpg", ".jpeg", ".webp", ".bmp", ".gif")) or "image/" in s

def _is_audio(name_or_type, url=""):
    s = f"{name_or_type or ''} {url or ''}".lower()
    return any(s.endswith(ext) for ext in (".mp3", ".webm", ".ogg", ".wav", ".m4a")) or "audio/" in s

def _is_video(name_or_type, url=""):
    s = f"{name_or_type or ''} {url or ''}".lower()
    return any(s.endswith(ext) for ext in (".mp4", ".mov", ".avi", ".mkv", ".flv")) or "video/" in s

def _resolve_file_on_disk(url_or_path, uploads_dir, temp_dir):
    """Faylni diskdan topadi (backend uploads, frontend public yoki base64)."""
    if not url_or_path or not isinstance(url_or_path, str):
        return None
    url_or_path = url_or_path.strip()
    if not url_or_path:
        return None

    # Base64 data URI bo'lsa
    if url_or_path.startswith("data:image/"):
        try:
            header, b64_data = url_or_path.split(";base64,", 1)
            raw = base64.b64decode(b64_data)
            ext = ".png"
            if "jpeg" in header or "jpg" in header:
                ext = ".jpg"
            elif "webp" in header:
                ext = ".webp"
            fname = f"b64_{uuid.uuid4().hex[:10]}{ext}"
            fpath = os.path.join(temp_dir, fname)
            with open(fpath, "wb") as bf:
                bf.write(raw)
            return fpath
        except Exception:
            return None

    # To'g'ridan-to'g'ri mavjud mutlaq yo'l
    clean = url_or_path.split("?")[0].split("#")[0].strip()
    if os.path.isabs(clean) and os.path.exists(clean):
        return os.path.abspath(clean)

    # Server URL prefiksini olib tashlash
    if "://" in clean:
        clean = "/" + clean.split("://", 1)[1].split("/", 1)[-1]

    filename = os.path.basename(clean)

    candidates = []
    if uploads_dir:
        candidates.append(os.path.join(uploads_dir, "tasks", filename))
        candidates.append(os.path.join(uploads_dir, filename))
        rel_sub = clean.lstrip("/\\")
        if rel_sub.startswith("uploads"):
            rel_sub = rel_sub[len("uploads"):].lstrip("/\\")
        candidates.append(os.path.join(uploads_dir, rel_sub))

        # Frontend public/uploads/tasks
        backend_dir = os.path.dirname(os.path.abspath(uploads_dir))
        frontend_tasks = os.path.join(backend_dir, "..", "frontend", "public", "uploads", "tasks", filename)
        candidates.append(os.path.abspath(frontend_tasks))

    for c in candidates:
        if os.path.exists(c) and os.path.isfile(c):
            return os.path.abspath(c)
    return None

def _resolve_public_url(url_or_path, server_url):
    """Excel ichidan bosilganda brauzerda ochiladigan to'liq URL."""
    if not url_or_path or not isinstance(url_or_path, str):
        return None
    url_or_path = url_or_path.strip()
    if not url_or_path or url_or_path.startswith("data:"):
        return None
    if url_or_path.startswith(("http://", "https://")):
        return url_or_path
    clean = "/" + url_or_path.lstrip("/")
    base = (server_url or "http://localhost:5000").rstrip("/")
    return f"{base}{clean}"

def _make_openpyxl_image(disk_path, temp_dir, max_w=115, max_h=66):
    """Diskdagi rasmdan openpyxl Image obyektini yaratadi."""
    if not disk_path or not os.path.exists(disk_path):
        return None
    try:
        im = PILImage.open(disk_path)
        try:
            im = ImageOps.exif_transpose(im)
        except Exception:
            pass
        orig_w, orig_h = im.size
        if orig_w <= 0 or orig_h <= 0:
            return None
        ratio = min(max_w / orig_w, max_h / orig_h)
        new_w = max(1, int(orig_w * ratio))
        new_h = max(1, int(orig_h * ratio))
        im = im.resize((new_w, new_h), PILImage.Resampling.LANCZOS)

        thumb_name = f"thumb_{uuid.uuid4().hex[:10]}.png"
        thumb_path = os.path.join(temp_dir, thumb_name)
        im.save(thumb_path, format="PNG")

        return OpenpyxlImage(thumb_path)
    except Exception as e:
        print(f"Rasm qayta ishlashda xatolik ({disk_path}): {e}", file=sys.stderr)
        return None

# ==============================================================================
# 3-VARAQ: «VAZIFALAR» (Jadval, kataklar ichiga fotosuratlar va havolalar bilan)
# ==============================================================================

def _tasks_sheet(ws: Worksheet, tasks: list, author: dict, now: datetime,
                 uploads_dir: str, server_url: str, temp_dir: str) -> None:
    columns = [
        Column("ID", 6),
        Column("Vazifa", 30, text=True),
        Column("Tavsif (Topshiriq)", 36, text=True),
        Column("Ijrochi", 22, text=True),
        Column("Hudud", 20, text=True),
        Column("Lavozim", 18, text=True),
        Column("Bergan rahbar", 18, text=True),
        Column("Holat", 16),
        Column("Berilgan", 16),
        Column("Muddat", 12, DATE_FMT),
        Column("Bajarilgan", 16),
        Column("Muddatga rioya", 16),
        Column("Qayta ishlash", 16),
        Column("Rad etish sababi", 30, text=True),
        # RAHBAR TOPSHIRIG'I MATERIALLARI
        Column("Rahbar ovozli topshirig'i", 26, text=True),
        Column("Rahbar fotosurati", 24, text=True),
        Column("Rahbar materiallari (Fayllar)", 30, text=True),
        # XODIM IJROSI VA DALILLARI
        Column("Xodim hisoboti (Izoh)", 38, text=True),
        Column("Xodim ovozli hisoboti", 26, text=True),
        Column("Ijro dalili (Fotosurat)", 24, text=True),
        Column("Ijro dalillari (Fayl va video)", 30, text=True),
    ]

    ordered = sorted(tasks, key=lambda t: (t.get("assignee_name") or "").lower())
    rows = []
    states = []
    author_name = author.get("name", "Administrator") if author else "Administrator"

    # Birinchi bosqich: asosiy ma'lumotlar massivini tayyorlash
    for t in ordered:
        state = compute_task_state(t, now)
        states.append(state)
        created_at = parse_date(t.get("created_at"))
        due_date = parse_date(t.get("due_date"))
        updated_at = parse_date(t.get("updated_at"))
        completed_val = updated_at if t.get("status") == "completed" else None
        status_label = STATUS_LABELS.get(t.get("status"), t.get("status"))

        rows.append([
            t.get("id"),
            t.get("title", ""),
            t.get("description") or "",
            t.get("assignee_name") or NOT_APPLICABLE,
            t.get("assignee_district") or t.get("assignee_region") or NOT_APPLICABLE,
            t.get("assignee_position") or "Xodim",
            author_name,
            status_label,
            created_at,
            due_date,
            completed_val,
            state,
            "Qayta ishlovda" if t.get("rework_required") else (
                f"Qaytarilgan: {t.get('rework_count')} marta" if t.get("rework_count") else NOT_APPLICABLE
            ),
            t.get("rework_reason") or NOT_APPLICABLE,
            # Placeholder textlar — quyida boyitiladi
            NOT_APPLICABLE, # Rahbar ovozi
            NOT_APPLICABLE, # Rahbar foto
            NOT_APPLICABLE, # Rahbar fayllar
            t.get("completion_note") or NOT_APPLICABLE, # Xodim hisoboti
            NOT_APPLICABLE, # Xodim ovozi
            NOT_APPLICABLE, # Xodim foto
            NOT_APPLICABLE, # Xodim fayllar
        ])

    col = _data_sheet(ws, columns, rows, SOLID["blue"])

    # Ikkinchi bosqich: fotosuratlarni katakka joylash va giperhavolalarni o'rnatish
    col_lead_audio = col["Rahbar ovozli topshirig'i"]
    col_lead_img = col["Rahbar fotosurati"]
    col_lead_files = col["Rahbar materiallari (Fayllar)"]
    col_emp_note = col["Xodim hisoboti (Izoh)"]
    col_emp_audio = col["Xodim ovozli hisoboti"]
    col_emp_img = col["Ijro dalili (Fotosurat)"]
    col_emp_files = col["Ijro dalillari (Fayl va video)"]

    for r, (t, state) in enumerate(zip(ordered, states), 2):
        ws.cell(row=r, column=col["Vazifa"]).font = Font(size=10, bold=True, color=INK)
        status_label = STATUS_LABELS.get(t.get("status"), t.get("status"))
        if status_label in STATUS_TONES:
            _tone(ws.cell(row=r, column=col["Holat"]), STATUS_TONES[status_label], bold=True)
        deadline = ws.cell(row=r, column=col["Muddatga rioya"])
        if state in DEADLINE_TONES:
            _tone(deadline, DEADLINE_TONES[state], bold=(state == OVERDUE))
        else:
            deadline.font = Font(size=10, italic=True, color=MUTED)

        # Rahbar materiallarini tahlil qilish
        lead_audio = t.get("audio_url")
        lead_attachments = _parse_items(t.get("attachments"))
        lead_images = [a for a in lead_attachments if _is_image(a.get("name"), a.get("url"))]
        lead_files = [a for a in lead_attachments if not _is_image(a.get("name"), a.get("url")) and not _is_audio(a.get("name"), a.get("url"))]

        # Xodim ijro dalillarini tahlil qilish
        emp_audio = t.get("completion_audio")
        emp_attachments = _parse_items(t.get("completion_attachments"))
        emp_images = [a for a in emp_attachments if _is_image(a.get("name"), a.get("url"))]
        emp_files = [a for a in emp_attachments if not _is_image(a.get("name"), a.get("url")) and not _is_audio(a.get("name"), a.get("url"))]

        has_embedded_photo = False

        # 1. Rahbar ovozli topshirig'i
        c_la = ws.cell(row=r, column=col_lead_audio)
        if lead_audio:
            pub_la = _resolve_public_url(lead_audio, server_url)
            c_la.value = "▶ Tinglash (Audio topshiriq)"
            if pub_la:
                c_la.hyperlink = pub_la
            c_la.font = Font(size=10, bold=True, color=LINK_BLUE, underline="single")
            c_la.alignment = CENTER
        else:
            c_la.value = NOT_APPLICABLE
            c_la.font = Font(size=10, color=MUTED)

        # 2. Rahbar fotosurati (katak ichiga embed qilish)
        c_li = ws.cell(row=r, column=col_lead_img)
        if lead_images:
            first_img = lead_images[0]
            disk_path = _resolve_file_on_disk(first_img.get("url"), uploads_dir, temp_dir)
            img_obj = _make_openpyxl_image(disk_path, temp_dir, max_w=115, max_h=66)
            pub_img = _resolve_public_url(first_img.get("url"), server_url)
            if img_obj:
                ws.add_image(img_obj, f"{get_column_letter(col_lead_img)}{r}")
                has_embedded_photo = True
                label = "🔍 To'liq ochish" + (f" (+{len(lead_images)-1})" if len(lead_images) > 1 else "")
                c_li.value = label
                if pub_img:
                    c_li.hyperlink = pub_img
                c_li.font = Font(size=9, color=LINK_BLUE, underline="single")
                c_li.alignment = Alignment(horizontal="center", vertical="bottom")
            else:
                c_li.value = f"📷 {first_img.get('name', 'foto')}"
                if pub_img:
                    c_li.hyperlink = pub_img
                c_li.font = Font(size=10, color=LINK_BLUE, underline="single")
        else:
            c_li.value = NOT_APPLICABLE
            c_li.font = Font(size=10, color=MUTED)

        # 3. Rahbar materiallari (Fayllar)
        c_lf = ws.cell(row=r, column=col_lead_files)
        if lead_files:
            names = [f"📎 {f.get('name') or os.path.basename(f.get('url') or '')}" for f in lead_files]
            c_lf.value = ", ".join(names)
            first_pub = _resolve_public_url(lead_files[0].get("url"), server_url)
            if first_pub:
                c_lf.hyperlink = first_pub
                c_lf.font = Font(size=10, color=LINK_BLUE, underline="single")
            else:
                c_lf.font = BODY_FONT
            c_lf.alignment = TEXT
        else:
            c_lf.value = NOT_APPLICABLE
            c_lf.font = Font(size=10, color=MUTED)

        # 4. Xodim ovozli hisoboti
        c_ea = ws.cell(row=r, column=col_emp_audio)
        if emp_audio:
            pub_ea = _resolve_public_url(emp_audio, server_url)
            c_ea.value = "🎙️ Tinglash (Xodim ovozi)"
            if pub_ea:
                c_ea.hyperlink = pub_ea
            c_ea.font = Font(size=10, bold=True, color=LINK_BLUE, underline="single")
            c_ea.alignment = CENTER
        else:
            c_ea.value = NOT_APPLICABLE
            c_ea.font = Font(size=10, color=MUTED)

        # 5. Xodim ijro dalili (Fotosurat katak ichida)
        c_ei = ws.cell(row=r, column=col_emp_img)
        if emp_images:
            first_img = emp_images[0]
            disk_path = _resolve_file_on_disk(first_img.get("url"), uploads_dir, temp_dir)
            img_obj = _make_openpyxl_image(disk_path, temp_dir, max_w=115, max_h=66)
            pub_img = _resolve_public_url(first_img.get("url"), server_url)
            if img_obj:
                ws.add_image(img_obj, f"{get_column_letter(col_emp_img)}{r}")
                has_embedded_photo = True
                label = "🔍 Dalilni ko'rish" + (f" (+{len(emp_images)-1})" if len(emp_images) > 1 else "")
                c_ei.value = label
                if pub_img:
                    c_ei.hyperlink = pub_img
                c_ei.font = Font(size=9, color=LINK_BLUE, underline="single")
                c_ei.alignment = Alignment(horizontal="center", vertical="bottom")
            else:
                c_ei.value = f"📷 {first_img.get('name', 'dalil')}"
                if pub_img:
                    c_ei.hyperlink = pub_img
                c_ei.font = Font(size=10, color=LINK_BLUE, underline="single")
        else:
            c_ei.value = NOT_APPLICABLE
            c_ei.font = Font(size=10, color=MUTED)

        # 6. Xodim ijro dalillari (Fayl va videolar)
        c_ef = ws.cell(row=r, column=col_emp_files)
        if emp_files:
            file_labels = []
            for f in emp_files:
                fn = f.get("name") or os.path.basename(f.get("url") or "") or "fayl"
                prefix = "🎥" if _is_video(fn, f.get("url")) else "📎"
                file_labels.append(f"{prefix} {fn}")
            c_ef.value = ", ".join(file_labels)
            first_pub = _resolve_public_url(emp_files[0].get("url"), server_url)
            if first_pub:
                c_ef.hyperlink = first_pub
                c_ef.font = Font(size=10, color=LINK_BLUE, underline="single")
            else:
                c_ef.font = BODY_FONT
            c_ef.alignment = TEXT
        else:
            c_ef.value = NOT_APPLICABLE
            c_ef.font = Font(size=10, color=MUTED)

        # Fotosurat mavjud qatorlarga mos balandlik o'rnatish
        if has_embedded_photo:
            ws.row_dimensions[r].height = 76

# ==============================================================================
# 4-VARAQ: «IJRO DALILLARI» (Foto-dosye kartalari: Rahbar topshirig'i vs Xodim ijrosi)
# ==============================================================================

def _evidence_priority(t):
    has_dalil = bool(
        t.get("completion_note") or
        t.get("completion_audio") or
        t.get("completion_attachments") or
        t.get("audio_url") or
        t.get("attachments")
    )
    status_order = {"completed": 1, "submitted": 2, "in_progress": 3, "pending": 4}
    return (0 if has_dalil else 1, status_order.get(t.get("status"), 5), -(t.get("id") or 0))

def _evidence_sheet(ws: Worksheet, tasks: list, author: dict, now: datetime,
                    uploads_dir: str, server_url: str, temp_dir: str) -> None:
    ws.sheet_properties.tabColor = SOLID["purple"]
    ws.sheet_view.showGridLines = True

    # Ustunlar kengligi
    ws.column_dimensions["A"].width = 2
    ws.column_dimensions["B"].width = 16
    ws.column_dimensions["C"].width = 22
    ws.column_dimensions["D"].width = 22
    ws.column_dimensions["E"].width = 22
    ws.column_dimensions["F"].width = 22
    ws.column_dimensions["G"].width = 16
    ws.column_dimensions["H"].width = 22
    ws.column_dimensions["I"].width = 22
    ws.column_dimensions["J"].width = 22
    ws.column_dimensions["K"].width = 22
    ws.column_dimensions["L"].width = 2

    # Sarlavha banneri
    _block(ws, 1, 2, 11, "RAHBAR TOPSHIRIQLARI VA XODIMLARNING IJRO DALILLARI DOSYESI",
           Font(size=14, bold=True, color=WHITE), _fill(NAVY), CENTER)
    _block(ws, 2, 2, 11,
           f"Shakllantirilgan sana: {now:%d.%m.%Y %H:%M}  |  Jami topshiriqlar: {len(tasks)} ta  |  Rahbar: {author.get('name', 'Administrator')}",
           Font(size=10, color="D6DCE5"), _fill(NAVY), CENTER)
    ws.row_dimensions[1].height = 32
    ws.row_dimensions[2].height = 20
    ws.row_dimensions[3].height = 10

    sorted_tasks = sorted(tasks, key=_evidence_priority)
    current_row = 4

    for t in sorted_tasks:
        tid = t.get("id")
        title = t.get("title") or "Topshiriq"
        status_raw = t.get("status")
        status_label = STATUS_LABELS.get(status_raw, status_raw)
        assignee = t.get("assignee_name") or "Biriktirilmagan"
        district = t.get("assignee_district") or t.get("assignee_region") or ""
        created_d = parse_date(t.get("created_at"))
        due_d = parse_date(t.get("due_date"))

        created_str = created_d.strftime("%d.%m.%Y") if created_d else "—"
        due_str = due_d.strftime("%d.%m.%Y") if due_d else "Muddatsiz"

        # Media parse
        lead_audio = t.get("audio_url")
        lead_attachments = _parse_items(t.get("attachments"))
        lead_images = [a for a in lead_attachments if _is_image(a.get("name"), a.get("url"))]
        lead_files = [a for a in lead_attachments if not _is_image(a.get("name"), a.get("url")) and not _is_audio(a.get("name"), a.get("url"))]

        emp_note = t.get("completion_note") or ""
        emp_audio = t.get("completion_audio")
        emp_attachments = _parse_items(t.get("completion_attachments"))
        emp_images = [a for a in emp_attachments if _is_image(a.get("name"), a.get("url"))]
        emp_files = [a for a in emp_attachments if not _is_image(a.get("name"), a.get("url")) and not _is_audio(a.get("name"), a.get("url"))]

        # 1. Karta sarlavhasi
        hdr_text = f"Topshiriq #{tid}: {title.upper()}  |  Holati: {status_label}  |  Ijrochi: {assignee} ({district})  |  Muddat: {due_str}"
        _block(ws, current_row, 2, 11, hdr_text, Font(size=11, bold=True, color=WHITE),
               _fill("2F5597"), Alignment(horizontal="left", vertical="center", indent=1), border=GRID)
        ws.row_dimensions[current_row].height = 26

        # 2. Bo'lim sarlavhalari (Rahbar vs Xodim)
        _block(ws, current_row + 1, 2, 6, f"🔷 RAHBAR TOPSHIRIG'I VA TALABLARI (Berilgan: {created_str})",
               Font(size=10, bold=True, color="1F4E78"), _fill("DDEBF7"), CENTER, border=GRID)
        _block(ws, current_row + 1, 7, 11, f"✅ XODIMNING IJRO HISOBOTI VA DALILLARI",
               Font(size=10, bold=True, color="375623"), _fill("E2EFDA"), CENTER, border=GRID)
        ws.row_dimensions[current_row + 1].height = 22

        # 3. Matn tavsifi va hisobot izohi
        desc_text = "Topshiriq mazmuni:\n" + (t.get("description") or "Topshiriq tavsifi berilmagan")
        note_text = "Xodim ijro hisoboti (Izoh):\n" + (emp_note if emp_note else "Yozma hisobot hali kiritilmagan")

        _block(ws, current_row + 2, 2, 6, desc_text, BODY_FONT,
               align=Alignment(horizontal="left", vertical="top", wrap_text=True), border=GRID)
        _block(ws, current_row + 2, 7, 11, note_text, BODY_FONT,
               align=Alignment(horizontal="left", vertical="top", wrap_text=True), border=GRID)

        lines_desc = max(2, sum(math.ceil(max(len(p), 1) / 70) for p in desc_text.split("\n")))
        lines_note = max(2, sum(math.ceil(max(len(p), 1) / 70) for p in note_text.split("\n")))
        ws.row_dimensions[current_row + 2].height = min(120, max(42, 15 * max(lines_desc, lines_note)))

        # 4. Ovozli xabarlar qatori
        c_la = _block(ws, current_row + 3, 2, 6, "", BODY_FONT, border=GRID, align=CENTER)
        if lead_audio:
            pub_la = _resolve_public_url(lead_audio, server_url)
            c_la.value = "▶ Tinglash: Rahbar ovozli topshirig'i"
            if pub_la:
                c_la.hyperlink = pub_la
            c_la.font = Font(size=10, bold=True, color=LINK_BLUE, underline="single")
        else:
            c_la.value = "🎙️ Rahbar ovozli topshirig'i yo'q"
            c_la.font = Font(size=10, color=MUTED, italic=True)

        c_ea = _block(ws, current_row + 3, 7, 11, "", BODY_FONT, border=GRID, align=CENTER)
        if emp_audio:
            pub_ea = _resolve_public_url(emp_audio, server_url)
            c_ea.value = "🎙️ Tinglash: Xodim ovozli hisoboti"
            if pub_ea:
                c_ea.hyperlink = pub_ea
            c_ea.font = Font(size=10, bold=True, color=LINK_BLUE, underline="single")
        else:
            c_ea.value = "🎙️ Xodim ovozli hisoboti yo'q"
            c_ea.font = Font(size=10, color=MUTED, italic=True)
        ws.row_dimensions[current_row + 3].height = 24

        # 5. Fotosuratlar qatori (Katta miniatyuralar katak ichida!)
        c_li = _block(ws, current_row + 4, 2, 6, "", BODY_FONT, border=GRID, align=CENTER)
        c_ei = _block(ws, current_row + 4, 7, 11, "", BODY_FONT, border=GRID, align=CENTER)

        has_card_photo = False

        if lead_images:
            first_img = lead_images[0]
            disk_path = _resolve_file_on_disk(first_img.get("url"), uploads_dir, temp_dir)
            img_obj = _make_openpyxl_image(disk_path, temp_dir, max_w=200, max_h=110)
            pub_img = _resolve_public_url(first_img.get("url"), server_url)
            if img_obj:
                ws.add_image(img_obj, f"B{current_row + 4}")
                has_card_photo = True
                c_li.value = "🔍 Rasmni to'liq ochish" + (f" (+{len(lead_images)-1})" if len(lead_images) > 1 else "")
                if pub_img:
                    c_li.hyperlink = pub_img
                c_li.font = Font(size=9, color=LINK_BLUE, underline="single")
                c_li.alignment = Alignment(horizontal="center", vertical="bottom")
            else:
                c_li.value = f"📷 [Foto: {first_img.get('name', 'rasm')}]"
                if pub_img:
                    c_li.hyperlink = pub_img
                c_li.font = Font(size=10, color=LINK_BLUE, underline="single")
        else:
            c_li.value = "📷 Rahbar fotosurati biriktirilmagan"
            c_li.font = Font(size=10, color=MUTED, italic=True)

        if emp_images:
            first_img = emp_images[0]
            disk_path = _resolve_file_on_disk(first_img.get("url"), uploads_dir, temp_dir)
            img_obj = _make_openpyxl_image(disk_path, temp_dir, max_w=200, max_h=110)
            pub_img = _resolve_public_url(first_img.get("url"), server_url)
            if img_obj:
                ws.add_image(img_obj, f"G{current_row + 4}")
                has_card_photo = True
                c_ei.value = "🔍 Ijro dalilini to'liq ochish" + (f" (+{len(emp_images)-1})" if len(emp_images) > 1 else "")
                if pub_img:
                    c_ei.hyperlink = pub_img
                c_ei.font = Font(size=9, color=LINK_BLUE, underline="single")
                c_ei.alignment = Alignment(horizontal="center", vertical="bottom")
            else:
                c_ei.value = f"📷 [Ijro dalili: {first_img.get('name', 'foto')}]"
                if pub_img:
                    c_ei.hyperlink = pub_img
                c_ei.font = Font(size=10, color=LINK_BLUE, underline="single")
        else:
            c_ei.value = "📷 Ijro fotosurati biriktirilmagan"
            c_ei.font = Font(size=10, color=MUTED, italic=True)

        ws.row_dimensions[current_row + 4].height = 105 if has_card_photo else 26

        # 6. Fayllar va hujjatlar qatori
        c_lf = _block(ws, current_row + 5, 2, 6, "", BODY_FONT, border=GRID, align=TEXT)
        if lead_files:
            file_names = [f"📎 {f.get('name') or os.path.basename(f.get('url') or '')}" for f in lead_files]
            c_lf.value = "Topshiriq materiallari: " + ", ".join(file_names)
            first_pub = _resolve_public_url(lead_files[0].get("url"), server_url)
            if first_pub:
                c_lf.hyperlink = first_pub
                c_lf.font = Font(size=10, color=LINK_BLUE, underline="single")
        else:
            c_lf.value = "Topshiriq materiallari: —"
            c_lf.font = Font(size=10, color=MUTED, italic=True)

        c_ef = _block(ws, current_row + 5, 7, 11, "", BODY_FONT, border=GRID, align=TEXT)
        if emp_files:
            file_names = []
            for f in emp_files:
                fn = f.get("name") or os.path.basename(f.get("url") or "") or "fayl"
                prefix = "🎥" if _is_video(fn, f.get("url")) else "📎"
                file_names.append(f"{prefix} {fn}")
            c_ef.value = "Ijro hujjatlari / dalillar: " + ", ".join(file_names)
            first_pub = _resolve_public_url(emp_files[0].get("url"), server_url)
            if first_pub:
                c_ef.hyperlink = first_pub
                c_ef.font = Font(size=10, color=LINK_BLUE, underline="single")
        else:
            c_ef.value = "Ijro hujjatlari: —"
            c_ef.font = Font(size=10, color=MUTED, italic=True)

        ws.row_dimensions[current_row + 5].height = 24

        # 7. Kartalar orasidagi bo'shliq
        ws.row_dimensions[current_row + 6].height = 12
        current_row += 7

    _print_setup(ws, landscape=True)

# ==============================================================================
# ASOSIY GENERATOR FUNKSIYASI
# ==============================================================================

def generate_report(data: dict, output_path: str):
    now = datetime.now()
    author = data.get("author", {})
    users = data.get("users", [])
    tasks = data.get("tasks", [])
    uploads_dir = data.get("uploadsDir") or ""
    server_url = data.get("serverUrl") or "http://localhost:5000"

    # Statistikalarni hisoblash
    total = UserStats()
    by_user = {u["id"]: UserStats() for u in users}

    for t in tasks:
        total.add(t, now)
        aid = t.get("assigned_to")
        if aid in by_user:
            by_user[aid].add(t, now)

    ranked = sorted(
        ((u, by_user[u["id"]]) for u in users),
        key=lambda p: (p[1].score is None, -(p[1].score or 0), -p[1].completed, p[0].get("name", "").lower()),
    )

    temp_dir = tempfile.mkdtemp(prefix="report_media_")
    try:
        wb = Workbook()
        _summary_sheet(wb.active, total, ranked, users, author, now)
        _employees_sheet(wb.create_sheet("Xodimlar"), ranked)
        _tasks_sheet(wb.create_sheet("Vazifalar"), tasks, author, now, uploads_dir, server_url, temp_dir)
        _evidence_sheet(wb.create_sheet("Ijro dalillari"), tasks, author, now, uploads_dir, server_url, temp_dir)

        wb.save(output_path)
    finally:
        shutil.rmtree(temp_dir, ignore_errors=True)

if __name__ == "__main__":
    if len(sys.argv) < 3:
        print("Usage: python report_generator.py <input_json_path> <output_xlsx_path>")
        sys.exit(1)

    input_json = sys.argv[1]
    output_xlsx = sys.argv[2]

    with open(input_json, "r", encoding="utf-8") as f:
        data = json.load(f)

    generate_report(data, output_xlsx)
    print(f"SUCCESS: {output_xlsx}")
